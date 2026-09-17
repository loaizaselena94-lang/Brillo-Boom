# servicios/reportes.py — Brillo-Boom
# Generación de reportes en PDF

from fpdf import FPDF
from datetime import datetime
from models import Usuario, Producto, Pedido

def generar_reporte_usuarios(usuarios):
    """Genera un PDF con listado de usuarios"""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'REPORTE DE USUARIOS - BRILLO-BOOM', ln=True, align='C')
    
    pdf.set_font('Arial', '', 10)
    pdf.cell(0, 10, f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}', ln=True, align='R')
    pdf.ln(5)
    
    # Encabezados
    pdf.set_font('Arial', 'B', 10)
    pdf.set_fill_color(228, 60, 126)  # Rosa brand
    pdf.set_text_color(255, 255, 255)
    pdf.cell(40, 8, 'ID', border=1, fill=True)
    pdf.cell(50, 8, 'Nombre', border=1, fill=True)
    pdf.cell(60, 8, 'Email', border=1, fill=True)
    pdf.cell(40, 8, 'Tipo', border=1, fill=True, ln=True)
    
    # Datos
    pdf.set_font('Arial', '', 9)
    pdf.set_text_color(0, 0, 0)
    for u in usuarios:
        tipo = 'Admin' if u.es_admin() else 'Cliente'
        pdf.cell(40, 7, str(u.id_usuario), border=1)
        pdf.cell(50, 7, u.nombre[:20], border=1)
        pdf.cell(60, 7, u.email[:25], border=1)
        pdf.cell(40, 7, tipo, border=1, ln=True)
    
    pdf.ln(5)
    pdf.set_font('Arial', '', 9)
    pdf.cell(0, 10, f'Total de usuarios: {len(usuarios)}', ln=True)
    
    return bytes(pdf.output())

def generar_reporte_productos(productos):
    """Genera un PDF con listado de productos"""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'REPORTE DE PRODUCTOS - BRILLO-BOOM', ln=True, align='C')
    
    pdf.set_font('Arial', '', 10)
    pdf.cell(0, 10, f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}', ln=True, align='R')
    pdf.ln(5)
    
    # Encabezados
    pdf.set_font('Arial', 'B', 9)
    pdf.set_fill_color(228, 60, 126)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(30, 8, 'ID', border=1, fill=True)
    pdf.cell(50, 8, 'Nombre', border=1, fill=True)
    pdf.cell(30, 8, 'Precio', border=1, fill=True)
    pdf.cell(25, 8, 'Stock', border=1, fill=True)
    pdf.cell(35, 8, 'Marca', border=1, fill=True, ln=True)
    
    # Datos
    pdf.set_font('Arial', '', 8)
    pdf.set_text_color(0, 0, 0)
    total_inventario = 0
    
    for p in productos:
        pdf.cell(30, 7, str(p.id_producto), border=1)
        pdf.cell(50, 7, p.nombre[:20], border=1)
        pdf.cell(30, 7, f'${p.precio:,.0f}', border=1)
        pdf.cell(25, 7, str(p.stock), border=1)
        pdf.cell(35, 7, p.empresa.nombre[:15], border=1, ln=True)
        total_inventario += p.precio * p.stock
    
    pdf.ln(5)
    pdf.set_font('Arial', '', 9)
    pdf.cell(0, 10, f'Total de productos: {len(productos)}', ln=True)
    pdf.cell(0, 10, f'Valor total del inventario: ${total_inventario:,.0f}', ln=True)
    
    return bytes(pdf.output())

def generar_reporte_pedidos(pedidos):
    """Genera un PDF con listado de pedidos"""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'REPORTE DE PEDIDOS - BRILLO-BOOM', ln=True, align='C')
    
    pdf.set_font('Arial', '', 10)
    pdf.cell(0, 10, f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}', ln=True, align='R')
    pdf.ln(5)
    
    # Encabezados
    pdf.set_font('Arial', 'B', 8)
    pdf.set_fill_color(228, 60, 126)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(20, 8, 'ID', border=1, fill=True)
    pdf.cell(35, 8, 'Usuario', border=1, fill=True)
    pdf.cell(35, 8, 'Producto', border=1, fill=True)
    pdf.cell(25, 8, 'Cantidad', border=1, fill=True)
    pdf.cell(30, 8, 'Valor', border=1, fill=True)
    pdf.cell(35, 8, 'Fecha', border=1, fill=True, ln=True)
    
    # Datos
    pdf.set_font('Arial', '', 7)
    pdf.set_text_color(0, 0, 0)
    total_ventas = 0
    
    for ped in pedidos:
        pdf.cell(20, 6, str(ped.id_pedido), border=1)
        pdf.cell(35, 6, ped.usuario.nombre[:15], border=1)
        pdf.cell(35, 6, ped.producto.nombre[:15], border=1)
        pdf.cell(25, 6, str(ped.cantidad), border=1)
        pdf.cell(30, 6, f'${ped.valor:,.0f}', border=1)
        fecha_str = ped.fecha.strftime("%d/%m/%Y") if hasattr(ped.fecha, 'strftime') else str(ped.fecha)[:10]
        pdf.cell(35, 6, fecha_str, border=1, ln=True)
        total_ventas += ped.valor
    
    pdf.ln(5)
    pdf.set_font('Arial', '', 9)
    pdf.cell(0, 10, f'Total de pedidos: {len(pedidos)}', ln=True)
    pdf.cell(0, 10, f'Ingresos totales: ${total_ventas:,.0f}', ln=True)
    
    return bytes(pdf.output())

def guardar_reporte_bd(tipo, archivo_bytes, nombre, id_admin):
    """Guarda el reporte en la base de datos"""
    from models import db, Reporte
    reporte = Reporte(
        tipo=tipo,
        archivo=archivo_bytes,
        nombre_archivo=nombre,
        id_usuario_admin=id_admin
    )
    db.session.add(reporte)
    db.session.commit()
    return reporte.id_reporte

def generar_reporte_productos_excel(productos):
    """Genera un archivo Excel (.xlsx) con listado de productos"""
    from openpyxl import Workbook
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"

    ws.append(['ID', 'Nombre', 'Precio', 'Stock', 'Marca'])
    for celda in ws[1]:
        celda.font = celda.font.copy(bold=True)

    total_inventario = 0
    for p in productos:
        ws.append([p.id_producto, p.nombre, p.precio, p.stock, p.empresa.nombre])
        total_inventario += p.precio * p.stock

    ws.append([])
    ws.append(['Total de productos', len(productos)])
    ws.append(['Valor total del inventario', total_inventario])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()